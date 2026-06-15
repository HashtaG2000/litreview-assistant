import streamlit as st
import os
import uuid
import tempfile
import re
from docx import Document as DocxDocument
from io import BytesIO
from ingestor import load_and_split, store_chunks
from retriever import search_similar, search_similar_filtered, check_relevance, check_counter_relevance
from llm import get_llm, call_llm_safe
from evaluator import (
    get_top_chunks_by_similarity,
    compute_review_quality,
    confidence_color,
)
from synthesizer import build_synthesis_matrix, detect_contradictions
from semantic_scholar import SemanticScholarClient

st.set_page_config(
    page_title="LitMap",
    page_icon="🗺️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');

    * { box-sizing: border-box; }

    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
        background-color: #0f1117;
        margin: 0;
        padding: 0;
    }

    .block-container {
        padding: 0 2rem 2rem 2rem !important;
        max-width: 100% !important;
    }

    /* ── Navbar ── */
    .navbar {
        display: flex;
        align-items: center;
        justify-content: space-between;
        padding: 18px 32px;
        border-bottom: 1px solid #1f2937;
        margin-bottom: 0;
        background: #0f1117;
        position: sticky;
        top: 0;
        z-index: 100;
    }
    .navbar-logo {
        font-size: 1.4rem;
        font-weight: 700;
        background: linear-gradient(135deg, #6366f1, #8b5cf6);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }
    .navbar-center { font-size: 0.95rem; color: #6b7280; }
    .navbar-right  { width: 120px; }

    /* ── Welcome ── */
    .hero { text-align: center; padding: 48px 16px 32px 16px; }
    .hero-title {
        font-size: 3rem; font-weight: 700;
        background: linear-gradient(135deg, #6366f1, #8b5cf6, #06b6d4);
        -webkit-background-clip: text; -webkit-text-fill-color: transparent;
        margin-bottom: 8px; line-height: 1.2;
    }
    .hero-sub { font-size: 1rem; color: #9ca3af; margin-bottom: 32px; }
    .chat-bubble {
        background: linear-gradient(135deg, #1e1b4b, #1e3a5f);
        border: 1px solid #4f46e5; border-radius: 16px;
        padding: 20px 28px; margin: 0 auto 32px auto; max-width: 620px; text-align: left;
    }
    .chat-bubble p { color: #e2e8f0; font-size: 0.95rem; line-height: 1.7; margin: 0 0 8px 0; }
    .chat-bubble p:last-child { margin-bottom: 0; }
    .form-label { font-size: 1rem; font-weight: 600; color: #e2e8f0; margin-bottom: 4px; }
    .form-sub   { font-size: 0.82rem; color: #6b7280; margin-bottom: 12px; }

    /* ── Understand screen ── */
    .understand-hero { text-align: center; padding: 40px 16px 24px 16px; }
    .understand-hero-title {
        font-size: 2.2rem; font-weight: 700;
        background: linear-gradient(135deg, #6366f1, #8b5cf6, #06b6d4);
        -webkit-background-clip: text; -webkit-text-fill-color: transparent;
        margin-bottom: 8px;
    }
    .understand-hero-sub { font-size: 0.95rem; color: #9ca3af; }
    .understand-card {
        background: #111827; border: 1px solid #374151;
        border-radius: 16px; padding: 32px 36px;
        margin: 24px auto 0 auto; max-width: 760px;
    }
    .understand-section {
        font-size: 0.72rem; font-weight: 600; letter-spacing: 0.1em;
        text-transform: uppercase; color: #6366f1; margin-bottom: 10px;
    }
    .understand-text { color: #d1d5db; font-size: 0.95rem; line-height: 1.85; }
    .paper-badge {
        display: inline-flex; align-items: center; gap: 6px;
        background: #1e1b4b; border: 1px solid #4f46e5; color: #818cf8;
        padding: 5px 14px; border-radius: 20px; font-size: 0.78rem;
        font-weight: 500; margin-bottom: 20px;
    }
    .understand-divider { border: none; border-top: 1px solid #1f2937; margin: 20px 0; }

    /* ── Validate ── */
    .col-label {
        font-size: 0.85rem; font-weight: 600; letter-spacing: 0.05em;
        text-transform: uppercase; padding-bottom: 10px; margin-bottom: 14px;
    }
    .col-upload  { color: #818cf8; border-bottom: 2px solid #4f46e5; }
    .col-green   { color: #4ade80; border-bottom: 2px solid #16a34a; }
    .col-red     { color: #f87171; border-bottom: 2px solid #dc2626; }
    .paper-card {
        border-radius: 10px; padding: 10px 14px; margin-bottom: 8px;
        font-size: 0.83rem; display: flex; align-items: flex-start;
        gap: 8px; word-break: break-word;
    }
    .paper-card-green { background: #052e16; border: 1px solid #16a34a; color: #bbf7d0; }
    .paper-card-red   { background: #1c0a0a; border: 1px solid #dc2626; color: #fecaca; }
    .paper-card-reason {
        font-size: 0.75rem; color: #9ca3af; margin-top: 4px;
        font-style: italic; line-height: 1.5;
    }
    .empty-state { color: #374151; font-size: 0.82rem; font-style: italic; padding: 12px 0; }

    /* ── Pending review cards ── */
    .pending-card {
        background: #111827; border: 1px solid #374151;
        border-radius: 12px; padding: 18px 20px; margin-bottom: 14px;
    }
    .pending-paper-name {
        font-size: 0.92rem; font-weight: 600; color: #e2e8f0;
        margin-bottom: 10px; word-break: break-word;
    }
    .verdict-badge-rel {
        display: inline-block; background: #052e16; border: 1px solid #16a34a;
        color: #4ade80; padding: 3px 12px; border-radius: 12px;
        font-size: 0.73rem; font-weight: 700; letter-spacing: 0.04em; margin-bottom: 10px;
    }
    .verdict-badge-norel {
        display: inline-block; background: #1c0a0a; border: 1px solid #dc2626;
        color: #f87171; padding: 3px 12px; border-radius: 12px;
        font-size: 0.73rem; font-weight: 700; letter-spacing: 0.04em; margin-bottom: 10px;
    }
    .confidence-badge {
        display: inline-block; background: #1a1a2e; border: 1px solid #374151;
        color: #9ca3af; padding: 3px 10px; border-radius: 12px;
        font-size: 0.73rem; font-weight: 600; margin-left: 8px; margin-bottom: 10px;
    }
    .pending-reason {
        color: #9ca3af; font-size: 0.84rem; line-height: 1.65; margin-bottom: 14px;
        padding: 10px 14px; background: #0f1117;
        border-left: 3px solid #374151; border-radius: 0 6px 6px 0;
    }
    .pending-section-header {
        font-size: 0.85rem; font-weight: 600; letter-spacing: 0.05em;
        text-transform: uppercase; color: #fbbf24; border-bottom: 2px solid #d97706;
        padding-bottom: 10px; margin: 24px 0 16px 0;
    }

    /* ── Suggestion cards (Semantic Scholar) ── */
    .suggestion-card {
        background: #0c1019; border: 1px solid #1e3a5f;
        border-radius: 10px; padding: 14px 18px; margin-bottom: 10px;
    }
    .suggestion-title { font-size: 0.88rem; font-weight: 600; color: #93c5fd; margin-bottom: 4px; }
    .suggestion-year { font-size: 0.75rem; color: #6b7280; margin-bottom: 6px; }
    .suggestion-relevance { font-size: 0.82rem; color: #a5b4fc; font-style: italic; margin-bottom: 6px; }
    .suggestion-abstract { font-size: 0.78rem; color: #6b7280; line-height: 1.5; }

    /* ── Review ── */
    .review-wrapper {
        background: #111827; border: 1px solid #1f2937;
        border-radius: 16px; padding: 40px 48px; margin-top: 24px;
    }
    .review-doc-title { font-size: 1.6rem; font-weight: 700; color: #f9fafb; margin-bottom: 4px; }
    .review-section {
        font-size: 0.75rem; font-weight: 600; letter-spacing: 0.1em;
        text-transform: uppercase; color: #6366f1; margin: 28px 0 10px 0;
    }
    .review-body { color: #d1d5db; font-size: 0.93rem; line-height: 1.85; }
    .ref-item { color: #9ca3af; font-size: 0.88rem; margin-bottom: 6px; }
    .hr { border: none; border-top: 1px solid #1f2937; margin: 28px 0; }

    /* ── Synthesis Matrix ── */
    .matrix-panel {
        background: #0c1019; border: 1px solid #1f2937;
        border-radius: 12px; padding: 24px 28px; margin-top: 28px;
    }
    .matrix-title {
        font-size: 0.72rem; font-weight: 600; letter-spacing: 0.1em;
        text-transform: uppercase; color: #06b6d4; margin-bottom: 16px;
    }
    .matrix-table { width: 100%; border-collapse: collapse; font-size: 0.80rem; }
    .matrix-table th {
        color: #6b7280; font-weight: 600; text-align: left;
        padding: 7px 12px; border-bottom: 1px solid #1f2937; white-space: nowrap;
    }
    .matrix-table td {
        color: #d1d5db; padding: 9px 12px; border-bottom: 1px solid #111827;
        vertical-align: top; line-height: 1.5;
    }
    .matrix-table tr:last-child td { border-bottom: none; }
    .matrix-paper-name { color: #818cf8; font-weight: 600; word-break: break-word; max-width: 180px; }

    /* ── Contradictions panel ── */
    .debate-panel {
        background: #0c1019; border: 1px solid #1f2937;
        border-radius: 12px; padding: 24px 28px; margin-top: 28px;
    }
    .debate-title {
        font-size: 0.72rem; font-weight: 600; letter-spacing: 0.1em;
        text-transform: uppercase; color: #f59e0b; margin-bottom: 14px;
    }
    .debate-body { color: #d1d5db; font-size: 0.88rem; line-height: 1.75; white-space: pre-wrap; }

    /* ── Review Quality Panel ── */
    .quality-panel {
        background: #0c1019; border: 1px solid #1f2937;
        border-radius: 12px; padding: 24px 28px; margin-top: 28px;
    }
    .quality-title {
        font-size: 0.72rem; font-weight: 600; letter-spacing: 0.1em;
        text-transform: uppercase; color: #6366f1; margin-bottom: 14px;
    }
    .quality-overall { font-size: 1rem; font-weight: 700; margin-bottom: 18px; }
    .quality-table { width: 100%; border-collapse: collapse; font-size: 0.82rem; }
    .quality-table th {
        color: #6b7280; font-weight: 600; text-align: left;
        padding: 6px 12px; border-bottom: 1px solid #1f2937;
    }
    .quality-table td { color: #d1d5db; padding: 8px 12px; border-bottom: 1px solid #111827; }
    .quality-table tr:last-child td { border-bottom: none; }

    /* ── Buttons ── */
    .stButton > button {
        background: linear-gradient(135deg, #6366f1, #8b5cf6) !important;
        color: white !important; border: none !important;
        border-radius: 8px !important; padding: 10px 20px !important;
        font-weight: 500 !important; font-size: 0.9rem !important;
        width: 100%; transition: opacity 0.2s !important;
    }
    .stButton > button:hover { opacity: 0.88 !important; }
    .stDownloadButton > button {
        background: linear-gradient(135deg, #059669, #0d9488) !important;
        color: white !important; border: none !important;
        border-radius: 8px !important; font-weight: 500 !important; width: 100%;
    }

    #MainMenu { visibility: hidden; }
    footer     { visibility: hidden; }
    header     { visibility: hidden; }
</style>
""", unsafe_allow_html=True)

# ── Session state ──────────────────────────────────────────────────────────────

defaults = {
    "screen": "welcome",
    "session_id": str(uuid.uuid4()),
    "research_idea": "",
    "accepted_papers": [],
    "rejected_papers": [],
    "pending_papers": [],
    "root_paper_name": "",
    "literature_review": "",
    "review_title": "",
    "review_quality": None,
    "synthesis_matrix": None,
    "contradiction_analysis": "",
    "suggested_papers": [],
    "processed_candidates": set(),
    "context_understanding": "",
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v

# ── Helpers ────────────────────────────────────────────────────────────────────

def parse_verdict(full_response):
    text = full_response.strip()
    lines = text.split("\n")
    first_upper = lines[0].strip().upper()
    reason_lines = [l.strip() for l in lines[1:] if l.strip()]
    reason = " ".join(reason_lines) if reason_lines else text
    if "NOT RELEVANT" in first_upper:
        return "NOT RELEVANT", reason
    elif "RELEVANT" in first_upper:
        return "RELEVANT", reason
    verdict = "NOT RELEVANT" if "NOT RELEVANT" in text.upper() else "RELEVANT"
    return verdict, reason


def _inline_citations(review_text: str, papers: list) -> str:
    """Post-process review text to append [N] after paper name mentions."""
    citation_map = {}
    for i, paper in enumerate(papers, 1):
        base = re.sub(r'\.pdf$', '', paper, flags=re.IGNORECASE)
        base = base.replace('_', ' ').replace('-', ' ')
        citation_map[base.lower()[:40]] = (base, i)

    result = review_text
    for key, (base, num) in citation_map.items():
        pattern = re.compile(re.escape(base), re.IGNORECASE)
        result = pattern.sub(f"{base} [{num}]", result, count=1)
    return result


def generate_docx(title, review_text, papers, matrix=None):
    doc = DocxDocument()
    doc.add_heading(title, 0)

    doc.add_heading("Literature Review", level=1)
    cited_text = _inline_citations(review_text, papers)
    doc.add_paragraph(cited_text)

    # Synthesis matrix table
    if matrix:
        doc.add_heading("Synthesis Matrix", level=1)
        headers = ["Paper", "Sample Size", "Methodology", "Key Findings", "Limitations"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
        for row in matrix:
            cells = table.add_row().cells
            cells[0].text = row.get("paper_title", "")
            cells[1].text = row.get("sample_size", "")
            cells[2].text = row.get("methodology", "")
            cells[3].text = row.get("key_findings", "")
            cells[4].text = row.get("limitations", "")

    doc.add_heading("References", level=1)
    for i, p in enumerate(papers, 1):
        doc.add_paragraph(f"{i}. {p}")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def navbar(center_text=""):
    st.markdown(f"""
    <div class="navbar">
        <div class="navbar-logo">🗺️ LitMap</div>
        <div class="navbar-center">{center_text}</div>
        <div class="navbar-right"></div>
    </div>
    """, unsafe_allow_html=True)


def quality_bar_html(pct: int, color: str) -> str:
    width = max(4, pct)
    return (
        f'<div style="display:flex;align-items:center;gap:8px;">'
        f'<div style="width:{width}px;height:6px;background:linear-gradient(90deg,#6366f1,#8b5cf6);'
        f'border-radius:3px;"></div>'
        f'<span style="color:{color};font-weight:600;">{pct}%</span>'
        f'</div>'
    )

# ═══════════════════════════════════════════════════════════════════════════════
# SCREEN 1 — Welcome
# ═══════════════════════════════════════════════════════════════════════════════

if st.session_state.screen == "welcome":

    navbar()

    st.markdown("""
    <div class="hero">
        <div class="hero-title">🗺️ LitMap</div>
        <div class="hero-sub">Your intelligent literature review assistant</div>
        <div class="chat-bubble">
            <p>👋 Welcome to <strong>LitMap</strong>. I help you build structured, validated literature reviews from your research papers.</p>
            <p>Describe your research idea, upload your root paper, and I'll validate every candidate paper you add — then draft your full literature review automatically.</p>
        </div>
    </div>
    """, unsafe_allow_html=True)

    _, col, _ = st.columns([1, 2, 1])
    with col:
        st.markdown('<div class="form-label">Your Research Idea</div>', unsafe_allow_html=True)
        st.markdown('<div class="form-sub">Describe your core research topic in 2-3 sentences. Be specific.</div>', unsafe_allow_html=True)
        research_idea = st.text_area(
            label="research_idea_input",
            label_visibility="collapsed",
            height=110,
            placeholder="e.g. This research investigates gamification in assembly line training for workers with cognitive disabilities, focusing on engagement and task accuracy..."
        )

        st.markdown("<div style='height:20px'></div>", unsafe_allow_html=True)

        st.markdown('<div class="form-label">Root Paper</div>', unsafe_allow_html=True)
        st.markdown('<div class="form-sub">Upload the foundation paper your research is built on.</div>', unsafe_allow_html=True)
        root_file = st.file_uploader(
            label="root_upload", label_visibility="collapsed",
            type="pdf", key="root_upload"
        )

        st.markdown("<div style='height:20px'></div>", unsafe_allow_html=True)

        if st.button("Start Mapping →", use_container_width=True):
            if not research_idea.strip():
                st.error("Please enter your research idea.")
            elif not root_file:
                st.error("Please upload your root paper.")
            else:
                st.session_state.research_idea = research_idea
                tmp_path = None
                try:
                    with st.status("Processing root paper...", expanded=True) as status:
                        st.write("Parsing PDF content...")
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                            tmp.write(root_file.read())
                            tmp_path = tmp.name
                        chunks = load_and_split(tmp_path)
                        st.write(f"Identified {len(chunks)} sections and chunks...")
                        st.write("Building knowledge base...")
                        store_chunks(chunks, root_file.name, st.session_state.session_id)
                        status.update(label="Root paper processed!", state="complete")
                except Exception as e:
                    st.error(f"Failed to process root paper: {e}")
                    st.stop()
                finally:
                    if tmp_path and os.path.exists(tmp_path):
                        os.unlink(tmp_path)

                st.session_state.root_paper_name = root_file.name
                st.session_state.accepted_papers.append(root_file.name)
                st.session_state.screen = "understand"
                st.rerun()

# ═══════════════════════════════════════════════════════════════════════════════
# SCREEN 2 — Understand
# ═══════════════════════════════════════════════════════════════════════════════

elif st.session_state.screen == "understand":

    navbar("Research Context")

    if not st.session_state.context_understanding:

        try:
            with st.status("Analysing your research context...", expanded=True) as status:
                st.write("Reading your research idea...")
                llm = get_llm()
                st.write("Retrieving key excerpts from root paper...")
                root_chunks = search_similar(
                    st.session_state.research_idea, k=6,
                    session_id=st.session_state.session_id
                )
                context = "\n\n".join([c.page_content for c in root_chunks])
                st.write("Building research context summary...")
                response = call_llm_safe(llm, f"""You are a research assistant reviewing a user's research idea and their root academic paper.

Research idea provided by the user:
"{st.session_state.research_idea}"

Key excerpts from the root paper "{st.session_state.root_paper_name}":
{context}

Write a clear, specific research context summary in 4–5 sentences covering:
1. The core research focus and objective
2. Key themes, methods, or concepts found in the paper
3. The gap or problem this research addresses
4. What types of candidate papers would be considered relevant for the literature review

Write in academic prose. Be specific. Do not use bullet points.
""")
                understanding = response.content if hasattr(response, "content") else str(response)
                st.session_state.context_understanding = understanding.strip()
                status.update(label="Context ready!", state="complete")
        except Exception as e:
            st.error(f"Could not generate research context — check your GROQ_API_KEY. Error: {e}")
            st.stop()

        st.rerun()

    else:
        st.markdown("""
        <div class="understand-hero">
            <div class="understand-hero-title">Here's what I understood</div>
            <div class="understand-hero-sub">Review the research context before we start validating candidate papers</div>
        </div>
        """, unsafe_allow_html=True)

        _, col, _ = st.columns([1, 2, 1])
        with col:
            st.markdown(f"""
            <div class="understand-card">
                <div class="paper-badge">📄 {st.session_state.root_paper_name}</div>
                <div class="understand-section">Research Context Summary</div>
                <div class="understand-text">{st.session_state.context_understanding}</div>
                <hr class="understand-divider">
                <div class="understand-section">Research Idea (as entered)</div>
                <div class="understand-text" style="color:#9ca3af;font-style:italic;">
                    "{st.session_state.research_idea}"
                </div>
            </div>
            """, unsafe_allow_html=True)

            st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)

            if st.button("Looks right — Start Validating Papers →", use_container_width=True):
                st.session_state.screen = "validate"
                st.rerun()

# ═══════════════════════════════════════════════════════════════════════════════
# SCREEN 3 — Validate
# ═══════════════════════════════════════════════════════════════════════════════

elif st.session_state.screen == "validate":

    navbar("Paper Validation")

    st.markdown(f"""
    <div class="chat-bubble" style="max-width:100%;margin:24px 0 20px 0;">
        <p>Root paper <strong>{st.session_state.root_paper_name}</strong> is loaded.
        Upload candidate papers one at a time — LitMap will analyse each one and explain its reasoning.
        You decide the final call: <strong>accept</strong> or <strong>reject</strong> each paper.</p>
    </div>
    """, unsafe_allow_html=True)

    # ── Upload ────────────────────────────────────────────────────────────────
    up_col, _, _ = st.columns([1, 1, 1], gap="large")
    with up_col:
        st.markdown('<div class="col-label col-upload">📄 Upload Candidate Paper</div>', unsafe_allow_html=True)
        candidate_file = st.file_uploader(
            label="candidate_upload", label_visibility="collapsed",
            type="pdf", key="candidate_upload"
        )

        if candidate_file and candidate_file.name not in st.session_state.processed_candidates:
            tmp_path = None
            try:
                with st.status(f"Analysing {candidate_file.name}...", expanded=True) as status:
                    st.write("Parsing PDF content...")
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                        tmp.write(candidate_file.read())
                        tmp_path = tmp.name
                    chunks = load_and_split(tmp_path)

                    st.write("Ranking chunks by semantic relevance...")
                    top_chunks, confidence = get_top_chunks_by_similarity(
                        st.session_state.research_idea, chunks, k=6
                    )

                    st.write("Running LLM relevance check...")
                    try:
                        full_verdict_text = check_relevance(
                            st.session_state.research_idea,
                            candidate_file.name,
                            top_chunks
                        )
                    except Exception as e:
                        st.error(f"LLM relevance check failed: {e}")
                        st.stop()

                    st.write("Generating devil's advocate counter-argument...")
                    counter_reason = check_counter_relevance(
                        st.session_state.research_idea,
                        candidate_file.name,
                        top_chunks
                    )

                    status.update(label="Analysis complete!", state="complete")

            except Exception as e:
                st.error(f"Failed to parse paper: {e}")
                st.stop()
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    os.unlink(tmp_path)

            verdict, reason = parse_verdict(full_verdict_text)

            st.session_state.processed_candidates.add(candidate_file.name)
            st.session_state.pending_papers.append({
                "name": candidate_file.name,
                "verdict": verdict,
                "reason": reason,
                "counter_reason": counter_reason,
                "confidence": confidence,
                "chunks": chunks,
            })
            st.rerun()

    # ── Pending review ────────────────────────────────────────────────────────
    if st.session_state.pending_papers:
        st.markdown('<div class="pending-section-header">⏳ Awaiting Your Review</div>', unsafe_allow_html=True)

        pending_action = None
        pending_action_idx = None

        for i, paper in enumerate(st.session_state.pending_papers):
            verdict_badge = (
                '<span class="verdict-badge-rel">✅ AI says: Relevant</span>'
                if paper["verdict"] == "RELEVANT"
                else '<span class="verdict-badge-norel">❌ AI says: Not Relevant</span>'
            )
            conf = paper["confidence"]
            col = confidence_color(conf)
            conf_badge = f'<span class="confidence-badge" style="color:{col};">Confidence: {conf}%</span>'

            st.markdown(f"""
            <div class="pending-card">
                <div class="pending-paper-name">📄 {paper["name"]}</div>
                {verdict_badge}{conf_badge}
                <div class="pending-reason"><strong>Reasoning:</strong> {paper["reason"]}</div>
            </div>
            """, unsafe_allow_html=True)

            if paper.get("counter_reason"):
                with st.expander("🔍 Why you might NOT include this paper"):
                    st.markdown(
                        f'<div style="color:#fbbf24;font-size:0.85rem;line-height:1.65;">{paper["counter_reason"]}</div>',
                        unsafe_allow_html=True
                    )

            btn_c1, btn_c2, _ = st.columns([1, 1, 4])
            with btn_c1:
                if st.button("✅ Accept", key=f"accept_{i}"):
                    pending_action = "accept"
                    pending_action_idx = i
            with btn_c2:
                if st.button("❌ Reject", key=f"reject_{i}"):
                    pending_action = "reject"
                    pending_action_idx = i

        if pending_action == "accept" and pending_action_idx is not None:
            paper = st.session_state.pending_papers[pending_action_idx]
            store_chunks(paper["chunks"], paper["name"], st.session_state.session_id)
            st.session_state.accepted_papers.append(paper["name"])
            st.session_state.pending_papers.pop(pending_action_idx)
            st.toast(f"✅ Accepted: {paper['name']}", icon="✅")
            st.rerun()

        elif pending_action == "reject" and pending_action_idx is not None:
            paper = st.session_state.pending_papers[pending_action_idx]
            st.session_state.rejected_papers.append({"name": paper["name"], "reason": paper["reason"]})
            st.session_state.pending_papers.pop(pending_action_idx)
            st.toast(f"❌ Rejected: {paper['name']}", icon="❌")
            st.rerun()

    # ── Citation Network Suggestions ─────────────────────────────────────────
    if len(st.session_state.accepted_papers) >= 2:
        with st.expander("💡 Papers from Citation Network (Semantic Scholar)"):
            if not st.session_state.suggested_papers:
                with st.spinner("Fetching citation network..."):
                    try:
                        client = SemanticScholarClient()
                        llm = get_llm()
                        suggestions = client.suggest_papers(
                            st.session_state.root_paper_name,
                            st.session_state.accepted_papers,
                            st.session_state.research_idea,
                            llm
                        )
                        st.session_state.suggested_papers = suggestions
                    except Exception:
                        st.session_state.suggested_papers = []

            if st.session_state.suggested_papers:
                st.markdown(
                    '<div style="color:#6b7280;font-size:0.8rem;margin-bottom:12px;">'
                    'Papers from the citation network of your root paper. Download and upload to validate them.</div>',
                    unsafe_allow_html=True
                )
                for s in st.session_state.suggested_papers:
                    st.markdown(f"""
                    <div class="suggestion-card">
                        <div class="suggestion-title">📄 {s['title']}</div>
                        <div class="suggestion-year">{s['year']}</div>
                        <div class="suggestion-relevance">✨ {s['relevance_summary']}</div>
                        <div class="suggestion-abstract">{s['abstract']}</div>
                    </div>
                    """, unsafe_allow_html=True)
            else:
                st.info("Could not find this paper in Semantic Scholar, or no related papers were found.")

    # ── Accepted / Rejected ───────────────────────────────────────────────────
    st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)
    col_rel, col_rej = st.columns([1, 1], gap="large")

    with col_rel:
        st.markdown('<div class="col-label col-green">✅ Accepted Papers</div>', unsafe_allow_html=True)
        if st.session_state.accepted_papers:
            for p in st.session_state.accepted_papers:
                st.markdown(f'<div class="paper-card paper-card-green">📄 {p}</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="empty-state">No accepted papers yet</div>', unsafe_allow_html=True)

    with col_rej:
        st.markdown('<div class="col-label col-red">❌ Rejected Papers</div>', unsafe_allow_html=True)
        if st.session_state.rejected_papers:
            for p in st.session_state.rejected_papers:
                st.markdown(f"""
                <div class="paper-card paper-card-red">
                    <div>
                        <div>📄 {p["name"]}</div>
                        <div class="paper-card-reason">💬 {p["reason"]}</div>
                    </div>
                </div>
                """, unsafe_allow_html=True)
        else:
            st.markdown('<div class="empty-state">No rejected papers yet</div>', unsafe_allow_html=True)

    # ── Done ──────────────────────────────────────────────────────────────────
    st.markdown("<div style='height:20px'></div>", unsafe_allow_html=True)
    _, done_col, _ = st.columns([2, 1, 2])
    with done_col:
        if st.button("Done — Draft Literature Review →", use_container_width=True):
            if len(st.session_state.accepted_papers) < 2:
                st.error("You need at least 2 accepted papers (including the root paper).")
            elif st.session_state.pending_papers:
                st.warning("Please review all pending papers before continuing.")
            else:
                st.session_state.screen = "review"
                st.rerun()

# ═══════════════════════════════════════════════════════════════════════════════
# SCREEN 4 — Literature Review
# ═══════════════════════════════════════════════════════════════════════════════

elif st.session_state.screen == "review":

    navbar("Literature Review")

    if not st.session_state.literature_review:

        try:
            with st.status("Generating your literature review...", expanded=True) as status:
                llm = get_llm()
                sid = st.session_state.session_id

                st.write("Retrieving balanced excerpts per paper...")
                paper_sections = []
                for paper in st.session_state.accepted_papers:
                    chunks = search_similar_filtered(
                        st.session_state.research_idea, paper, k=4, session_id=sid
                    )
                    if chunks:
                        section_text = "\n\n".join([c.page_content for c in chunks])
                        paper_sections.append(f"[Source: {paper}]\n{section_text}")

                context = "\n\n---\n\n".join(paper_sections) if paper_sections else "No excerpts available."

                st.write("Generating title...")
                title_response = call_llm_safe(
                    llm,
                    f"Generate a concise academic title (maximum 12 words) for a literature review on this topic: "
                    f"{st.session_state.research_idea}. Return only the title, nothing else."
                )
                st.session_state.review_title = (
                    title_response.content.strip()
                    if hasattr(title_response, "content")
                    else str(title_response).strip()
                )

                st.write("Drafting literature review...")
                review_response = call_llm_safe(llm, f"""
You are an academic research assistant. Write a structured literature review.

Research idea: {st.session_state.research_idea}

Accepted papers: {', '.join(st.session_state.accepted_papers)}

Excerpts organised by paper (each section is labelled with its source):
{context}

Structure:
- Introduction paragraph summarising the research landscape
- Numbered analysis per paper (use the paper name in bold, 2-3 sentences linking it to the research idea)
- Conclusion paragraph synthesising the reviewed works

Every accepted paper must be discussed. Use clear academic language.
""")
                raw_review = (
                    review_response.content
                    if hasattr(review_response, "content")
                    else str(review_response)
                )
                st.session_state.literature_review = raw_review

                st.write("Extracting synthesis matrix...")
                try:
                    matrix = build_synthesis_matrix(
                        llm,
                        st.session_state.accepted_papers,
                        search_similar_filtered,
                        st.session_state.research_idea,
                        sid
                    )
                    st.session_state.synthesis_matrix = matrix
                except Exception:
                    st.session_state.synthesis_matrix = None

                st.write("Detecting contradictions and consensus...")
                try:
                    if st.session_state.synthesis_matrix and len(st.session_state.synthesis_matrix) >= 2:
                        st.session_state.contradiction_analysis = detect_contradictions(
                            llm, st.session_state.synthesis_matrix, st.session_state.research_idea
                        )
                except Exception:
                    st.session_state.contradiction_analysis = ""

                st.write("Computing retrieval quality metrics...")
                try:
                    def _search_fn(q, p, k):
                        return search_similar_filtered(q, p, k=k, session_id=sid)

                    paper_stats, overall_confidence = compute_review_quality(
                        st.session_state.research_idea,
                        st.session_state.accepted_papers,
                        _search_fn,
                        k_per_paper=4
                    )
                    st.session_state.review_quality = {
                        "paper_stats": paper_stats,
                        "overall": overall_confidence,
                    }
                except Exception:
                    st.session_state.review_quality = None

                status.update(label="Literature review ready!", state="complete")

        except Exception as e:
            st.error(f"Literature review generation failed: {e}. Please try again.")
            st.stop()

        st.rerun()

    else:
        review_html = st.session_state.literature_review.replace("\n", "<br>")

        st.markdown(f"""
        <div class="review-wrapper">
            <div class="review-doc-title">{st.session_state.review_title}</div>
            <div class="review-section">Literature Review</div>
            <div class="review-body">{review_html}</div>
            <div class="hr"></div>
            <div class="review-section">References</div>
            {"".join([f'<div class="ref-item">{i+1}. {p}</div>' for i, p in enumerate(st.session_state.accepted_papers)])}
        </div>
        """, unsafe_allow_html=True)

        # ── Synthesis Matrix ──────────────────────────────────────────────────
        matrix = st.session_state.synthesis_matrix
        if matrix:
            rows_html = ""
            for row in matrix:
                rows_html += f"""
                <tr>
                    <td class="matrix-paper-name">{row['paper_title']}</td>
                    <td>{row['sample_size']}</td>
                    <td>{row['methodology']}</td>
                    <td>{row['key_findings']}</td>
                    <td>{row['limitations']}</td>
                </tr>
                """
            st.markdown(f"""
            <div class="matrix-panel">
                <div class="matrix-title">📊 Synthesis Matrix</div>
                <table class="matrix-table">
                    <tr>
                        <th>Paper</th>
                        <th>Sample Size</th>
                        <th>Methodology</th>
                        <th>Key Findings</th>
                        <th>Limitations</th>
                    </tr>
                    {rows_html}
                </table>
            </div>
            """, unsafe_allow_html=True)

        # ── Contradictions & Consensus ────────────────────────────────────────
        if st.session_state.contradiction_analysis:
            debate_html = st.session_state.contradiction_analysis.replace("\n", "<br>")
            st.markdown(f"""
            <div class="debate-panel">
                <div class="debate-title">⚡ Debates & Consensus</div>
                <div class="debate-body">{debate_html}</div>
            </div>
            """, unsafe_allow_html=True)

        # ── Retrieval Quality ─────────────────────────────────────────────────
        quality = st.session_state.review_quality
        if quality:
            overall = quality["overall"]
            paper_stats = quality["paper_stats"]
            overall_col = confidence_color(overall)

            rows_html = ""
            for paper, stats in paper_stats.items():
                conf = stats["confidence"]
                col = confidence_color(conf)
                rows_html += f"""
                <tr>
                    <td style="word-break:break-word;max-width:320px;">{paper}</td>
                    <td style="text-align:center;">{stats['chunks_retrieved']}</td>
                    <td>{quality_bar_html(conf, col)}</td>
                </tr>
                """

            st.markdown(f"""
            <div class="quality-panel">
                <div class="quality-title">📈 Retrieval Quality</div>
                <div class="quality-overall" style="color:{overall_col};">
                    Overall Retrieval Confidence: {overall}%
                </div>
                <table class="quality-table">
                    <tr>
                        <th>Paper</th>
                        <th style="text-align:center;">Chunks Retrieved</th>
                        <th>Relevance Confidence</th>
                    </tr>
                    {rows_html}
                </table>
            </div>
            """, unsafe_allow_html=True)

        st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)

        dl_col, reset_col, _ = st.columns([1, 1, 2])

        with dl_col:
            docx_buf = generate_docx(
                st.session_state.review_title,
                st.session_state.literature_review,
                st.session_state.accepted_papers,
                matrix=st.session_state.synthesis_matrix,
            )
            st.download_button(
                label="⬇️ Download Word Document",
                data=docx_buf,
                file_name="LitMap_Literature_Review.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

        with reset_col:
            if st.button("← Start New Review", use_container_width=True):
                for key in list(st.session_state.keys()):
                    del st.session_state[key]
                st.rerun()
